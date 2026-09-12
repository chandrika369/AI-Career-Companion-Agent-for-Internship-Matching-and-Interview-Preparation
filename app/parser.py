import re
import os
from typing import Dict, List, Any, Optional

from pypdf import PdfReader
from docx import Document


# ============================================================
# TEXT EXTRACTION
# ============================================================

def extract_text(file_path: str, file_extension: str) -> str:
    """
    Extract text from PDF or DOCX files cleanly with structure preservation.
    """
    if file_extension.lower() == ".pdf":
        reader = PdfReader(file_path)
        pages = []
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                pages.append(page_text)
        return "\n".join(pages)

    elif file_extension.lower() == ".docx":
        doc = Document(file_path)
        lines = []

        # 1. Paragraphs
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if text:
                lines.append(text)

        # 2. Tables (with merged cell deduplication and line preservation)
        for table in doc.tables:
            for row in table.rows:
                seen_cell_texts = set()
                for cell in row.cells:
                    cell_text = cell.text.strip()
                    if cell_text and cell_text not in seen_cell_texts:
                        seen_cell_texts.add(cell_text)
                        for sub_line in cell_text.splitlines():
                            sub = sub_line.strip()
                            if sub:
                                lines.append(sub)

        return "\n".join(lines)

    else:
        raise ValueError("Unsupported file format")


# ============================================================
# GENERAL TEXT CLEANING
# ============================================================

def clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    # Normalize excessive spaces
    text = re.sub(r"[ \t]+", " ", text)
    # Normalize excessive blank lines
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# ============================================================
# BASIC CONTACT INFORMATION
# ============================================================

def extract_email(text: str) -> Optional[str]:
    match = re.search(
        r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b",
        text
    )
    return match.group(0) if match else None


def extract_phone(text: str) -> Optional[str]:
    patterns = [
        r"(?<!\d)\+?\d{1,3}[-.\s]?\(?\d{2,4}\)?[-.\s]?\d{3,4}[-.\s]?\d{3,4}(?!\d)",
        r"(?<!\d)\d{10}(?!\d)"
    ]
    for pattern in patterns:
        match = re.search(pattern, text)
        if match:
            return match.group(0).strip()
    return None


def extract_linkedin(text: str) -> Optional[str]:
    match = re.search(
        r"(https?://(?:www\.)?linkedin\.com/[^\s]+)",
        text,
        re.IGNORECASE
    )
    return match.group(1).rstrip(".,)") if match else None


def extract_github(text: str) -> Optional[str]:
    match = re.search(
        r"(https?://(?:www\.)?github\.com/[^\s]+)",
        text,
        re.IGNORECASE
    )
    return match.group(1).rstrip(".,)") if match else None


# ============================================================
# NAME EXTRACTION
# ============================================================

def extract_name(text: str) -> Optional[str]:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if not lines:
        return None

    # Look at first few lines
    for line in lines[:8]:
        if "@" in line:
            continue
        if re.search(r"https?://", line):
            continue
        if re.search(r"\d{5,}", line):
            continue
        if line.upper() in {"RESUME", "CV", "CURRICULUM VITAE", "PROFILE", "PORTFOLIO"}:
            continue
        if re.search(r"\b(course|mobile|email|cgpa|social|github|linkedin)\b", line, re.IGNORECASE):
            continue

        # Name is generally 1-6 words, under 50 chars
        if 1 <= len(line.split()) <= 6 and len(line) <= 50:
            if not re.search(
                r"\b(education|skills|experience|projects|summary|objective|certification|academic)\b",
                line,
                re.IGNORECASE
            ):
                return line
    return None


# ============================================================
# SECTION HEADING DETECTION
# ============================================================

SECTION_ALIASES = {
    "summary": [
        "summary", "professional summary", "profile", "profile summary",
        "career summary", "about me", "about", "objective", "career objective",
        "professional profile"
    ],
    "skills": [
        "skills", "skill set", "core skills", "key skills", "technical skills",
        "technical proficiency", "competencies", "core competencies", "expertise"
    ],
    "technical_skills": [
        "technical skills", "technical proficiency", "technical expertise",
        "technologies", "technology stack", "tech stack", "tools and technologies",
        "programming skills", "subjects / electives", "electives"
    ],
    "education": [
        "education", "educational background", "academic background",
        "academic qualifications", "qualifications", "academic details",
        "education qualifications", "academics"
    ],
    "work_experience": [
        "work experience", "professional experience", "employment history",
        "employment experience", "career history", "work history",
        "experience", "professional background"
    ],
    "internships": [
        "internship", "internships", "internship experience", "intern experiences",
        "virtual internship", "virtual internships", "internship details"
    ],
    "projects": [
        "projects", "project", "academic projects", "personal projects",
        "major projects", "key projects", "project experience", "selected projects"
    ],
    "certifications": [
        "certifications", "certificates", "certification", "professional certifications",
        "licenses and certifications", "licenses", "training and certifications"
    ],
    "languages": [
        "languages", "language", "languages known", "language skills"
    ],
    "achievements": [
        "achievements", "accomplishments", "awards", "honors", "honours",
        "recognition", "awards and recognitions", "awards & recognitions"
    ],
    "soft_skills": [
        "soft skills", "interpersonal skills", "personal skills", "behavioral skills"
    ],
    "responsibility": [
        "position of responsibility", "positions of responsibility", "leadership",
        "responsibilities", "extracurricular activities"
    ]
}

TABLE_HEADERS = {
    "course", "institute / college", "institute", "college", "board / university",
    "board", "university", "score", "year", "percentage", "cgpa", "qualification",
    "degree / diploma", "degree", "marks", "s.no", "sr. no", "sr no", "s no",
    "certification", "certifying authority", "authority", "date", "mode", "status"
}


def normalize_heading(text: str) -> str:
    text = text.strip()
    text = re.sub(r"[:\-#*|•]+$", "", text)
    # Collapse repeated consecutive words (e.g. 'INTERNSHIPS INTERNSHIPS' -> 'INTERNSHIPS')
    text = re.sub(r"\b(\w+)(?:\s+\1\b)+", r"\1", text, flags=re.IGNORECASE)
    text = re.sub(r"\s+", " ", text)
    return text.lower().strip()


def is_table_header(line: str) -> bool:
    cleaned = line.lower().strip(" :|-#*•")
    if cleaned in TABLE_HEADERS:
        return True
    words = [w.strip(" /|-") for w in re.split(r"[\s/|]+", cleaned) if w.strip(" /|-")]
    if words and all(w in {
        "course", "institute", "college", "board", "university", "score", "year",
        "percentage", "cgpa", "marks", "grade", "qualification", "degree", "sno",
        "no", "sr", "certification", "certifying", "authority"
    } for w in words):
        return True
    return False


def detect_section_heading(line: str) -> Optional[str]:
    normalized = normalize_heading(line)
    if not normalized or len(normalized) > 55:
        return None

    # Exact match
    for section, aliases in SECTION_ALIASES.items():
        for alias in aliases:
            if normalized == alias.lower():
                return section

    # Partial prefix / suffix match for short headings
    for section, aliases in SECTION_ALIASES.items():
        for alias in aliases:
            alias_norm = alias.lower()
            if len(normalized.split()) <= 4 and (
                normalized.startswith(alias_norm) or normalized.endswith(alias_norm)
            ):
                return section

    return None


# ============================================================
# SECTION EXTRACTION
# ============================================================

def extract_sections(text: str) -> Dict[str, List[str]]:
    sections: Dict[str, List[str]] = {"header": []}
    current_section = "header"

    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            continue

        detected = detect_section_heading(line)
        if detected:
            current_section = detected
            if current_section not in sections:
                sections[current_section] = []
            continue

        sections.setdefault(current_section, [])
        sections[current_section].append(line)

    return sections


def section_text(sections: Dict[str, List[str]], name: str) -> str:
    return "\n".join(sections.get(name, []))


# ============================================================
# SKILLS EXTRACTION
# ============================================================

def split_skill_items(text: str) -> List[str]:
    if not text:
        return []

    items = []
    parts = re.split(r"[,;|\n•●▪◦○*]", text)

    for item in parts:
        item = item.strip(" -–—:.\t\r\n")
        if not item:
            continue

        # Reject category labels
        if re.match(
            r"^(programming languages|libraries|tools|web technologies|concepts|"
            r"technologies|subjects\s*/\s*electives|technical proficiency|skills|core skills)\s*$",
            item,
            re.IGNORECASE
        ):
            continue

        # Reject section titles
        if re.match(r"^(internships?|projects?|certifications?|academic details|declaration|education)\b", item, re.IGNORECASE):
            continue

        # Reject long narrative descriptions or full sentences
        if len(item) > 55 or len(item.split()) > 7:
            continue
        if re.search(r"\b(successfully|completed|offered by|gained hands-on|awarded|predictive model|forecast|developed|implemented|utilizing|currently)\b", item, re.IGNORECASE):
            continue

        items.append(item)

    return list(dict.fromkeys(items))


def extract_skills(sections: Dict[str, List[str]]) -> List[str]:
    text = section_text(sections, "skills")
    return split_skill_items(text)


def extract_technical_skills(sections: Dict[str, List[str]]) -> Dict[str, List[str]]:
    text = section_text(sections, "technical_skills")
    result = {
        "programming_languages": [],
        "libraries_tools": [],
        "web_technologies": [],
        "concepts": []
    }

    if not text:
        return result

    current_category = None
    category_patterns = {
        "programming_languages": ["programming languages", "programming"],
        "libraries_tools": ["libraries/tools", "libraries", "tools"],
        "web_technologies": ["web technologies", "web development", "web"],
        "concepts": ["concepts", "core concepts", "subjects"]
    }

    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue

        normalized = line.lower().strip(": ")
        found_category = None

        for category, aliases in category_patterns.items():
            for alias in aliases:
                if normalized == alias:
                    found_category = category
                    break
            if found_category:
                break

        if found_category:
            current_category = found_category
            continue

        if current_category:
            result[current_category].extend(split_skill_items(line))
        else:
            result["libraries_tools"].extend(split_skill_items(line))

    for key in result:
        result[key] = list(dict.fromkeys(result[key]))

    return result


# ============================================================
# EDUCATION EXTRACTION
# ============================================================

def extract_education(sections: Dict[str, List[str]]) -> List[str]:
    edu_lines = sections.get("education", [])
    header_lines = sections.get("header", [])
    items = []

    # 1. Degree / Course from header (e.g. Course : B.Tech, Computer Science and Engineering...)
    for l in header_lines:
        if re.search(r"\b(course|degree|branch|b\.?tech|b\.?e\.?|m\.?tech|bca|mca|b\.?sc|m\.?sc|bachelor|master)\b", l, re.IGNORECASE):
            cleaned = re.sub(r"^(course|degree|branch)\s*[:\-]\s*", "", l, flags=re.IGNORECASE).strip()
            if cleaned and not is_table_header(cleaned):
                items.append(cleaned)

    # 2. Process education section lines, collapsing table cells into structured entries
    i = 0
    while i < len(edu_lines):
        line = edu_lines[i].strip()
        if not line or is_table_header(line):
            i += 1
            continue

        parts = [line]
        j = i + 1
        while j < len(edu_lines) and j <= i + 4:
            next_l = edu_lines[j].strip()
            if not next_l or is_table_header(next_l):
                j += 1
                continue
            if re.search(r"^(class\s+(?:x|xii|10|12)|b\.?tech|b\.?e\.?|m\.?tech|bachelor|master|diploma|secondary|intermediate)\b", next_l, re.IGNORECASE):
                break
            parts.append(next_l)
            j += 1

        combined = " - ".join(parts)
        if combined not in items:
            items.append(combined)
        i = j if j > i + 1 else i + 1

    return list(dict.fromkeys(items))


# ============================================================
# INTERNSHIPS EXTRACTION
# ============================================================

def extract_internships(sections: Dict[str, List[str]]) -> List[str]:
    lines = sections.get("internships", [])
    internships = []
    current_entry = []

    for l in lines:
        l = l.strip()
        if not l or is_table_header(l):
            continue
        if re.match(r"^(internships?|projects?|certifications?|academic)\b", l, re.IGNORECASE):
            continue

        # New internship detected if line is short and has dates or job titles
        if re.search(r"\b(202\d|developer|intern|engineer|lead|assistant|specialist)\b", l, re.IGNORECASE) and len(l) < 120 and (not current_entry or len(current_entry) >= 2):
            if current_entry:
                internships.append("\n".join(current_entry))
            current_entry = [l]
        else:
            if current_entry:
                current_entry.append(l)
            else:
                current_entry = [l]

    if current_entry:
        internships.append("\n".join(current_entry))

    return internships


# ============================================================
# PROJECTS EXTRACTION
# ============================================================

def extract_projects(sections: Dict[str, List[str]]) -> List[str]:
    lines = sections.get("projects", [])
    projects = []
    current_entry = []

    for l in lines:
        l = l.strip()
        if not l or is_table_header(l):
            continue
        if re.match(r"^(projects?|internships?|certifications?|academic)\b", l, re.IGNORECASE):
            continue

        # Project title indicators
        if re.search(r"\b(live|github|web application|deep learning|machine learning|app|project|system|model|tracker)\b", l, re.IGNORECASE) and len(l) < 120 and (not current_entry or len(current_entry) >= 2):
            if current_entry:
                projects.append("\n".join(current_entry))
            current_entry = [l]
        else:
            if current_entry:
                current_entry.append(l)
            else:
                current_entry = [l]

    if current_entry:
        projects.append("\n".join(current_entry))

    return projects


# ============================================================
# CERTIFICATIONS EXTRACTION
# ============================================================

def extract_certifications(sections: Dict[str, List[str]]) -> List[str]:
    lines = sections.get("certifications", [])
    certs = []
    i = 0

    while i < len(lines):
        l = lines[i].strip()
        if not l or is_table_header(l):
            i += 1
            continue

        # If next line is certifying authority, pair them up: 'React.js (Scaler)'
        if i + 1 < len(lines):
            next_l = lines[i + 1].strip()
            if next_l and not is_table_header(next_l) and len(next_l) <= 40 and not re.search(r"certificate|certification", next_l, re.IGNORECASE):
                certs.append(f"{l} ({next_l})")
                i += 2
                continue

        certs.append(l)
        i += 1

    return list(dict.fromkeys(certs))


# ============================================================
# ACHIEVEMENTS & LANGUAGES & OTHERS
# ============================================================

def extract_achievements(sections: Dict[str, List[str]]) -> List[str]:
    lines = sections.get("achievements", [])
    items = []
    for l in lines:
        l = l.strip()
        if l and not is_table_header(l) and not re.match(r"^(achievements?|awards?)\b", l, re.IGNORECASE):
            items.append(l)
    return list(dict.fromkeys(items))


def extract_languages(sections: Dict[str, List[str]]) -> List[str]:
    text = section_text(sections, "languages")
    return split_skill_items(text)


def extract_summary(sections: Dict[str, List[str]]) -> Optional[str]:
    text = section_text(sections, "summary")
    if not text:
        return None
    return " ".join(text.split())


def extract_address(text: str, sections: Dict[str, List[str]]) -> Optional[str]:
    header_lines = sections.get("header", [])
    email = extract_email(text)
    phone = extract_phone(text)

    for line in header_lines:
        line = line.strip()
        if not line or (email and email in line) or (phone and phone in line):
            continue
        if re.search(r"linkedin\.com|github\.com", line, re.IGNORECASE):
            continue
        if detect_section_heading(line):
            continue
        if re.search(r"\b(street|road|avenue|lane|apartment|building|city|district|state|country|pincode|zip|india)\b", line, re.IGNORECASE):
            cleaned = re.sub(r"\s{2,}", " ", line).strip(" ,-|:")
            if cleaned:
                return cleaned
    return None


def extract_other_information(sections: Dict[str, List[str]], original_text: str) -> List[str]:
    known_section_names = set(SECTION_ALIASES.keys())
    information = []

    header_lines = sections.get("header", [])
    for line in header_lines:
        line = line.strip()
        if not line:
            continue
        if extract_email(line) or extract_phone(line):
            continue
        if re.search(r"linkedin\.com|github\.com", line, re.IGNORECASE):
            continue
        if len(line.split()) <= 6 and len(line) <= 50 and not re.search(r"course|degree|cgpa|social", line, re.IGNORECASE):
            continue
        information.append(line)

    for section_name, lines in sections.items():
        if section_name in known_section_names or section_name == "header":
            continue
        for line in lines:
            if line.strip() and not is_table_header(line.strip()):
                information.append(line.strip())

    return list(dict.fromkeys(information))


# ============================================================
# MAIN RESUME PARSER
# ============================================================

def parse_resume(text: str) -> Dict[str, Any]:
    text = clean_text(text)
    sections = extract_sections(text)

    technical_skills = extract_technical_skills(sections)
    general_skills = extract_skills(sections)

    all_skills = []
    all_skills.extend(general_skills)
    for category in technical_skills.values():
        all_skills.extend(category)
    all_skills = list(dict.fromkeys(all_skills))

    internships = extract_internships(sections)
    work_exp = [internships[0]] if internships else []

    return {
        "full_name": extract_name(text),
        "email": extract_email(text),
        "phone": extract_phone(text),
        "address": extract_address(text, sections),
        "linkedin": extract_linkedin(text),
        "github": extract_github(text),
        "summary": extract_summary(sections),
        "skills": all_skills,
        "education": extract_education(sections),
        "work_experience": work_exp,
        "projects": extract_projects(sections),
        "certifications": extract_certifications(sections),
        "internships": internships,
        "languages": extract_languages(sections),
        "achievements": extract_achievements(sections),
        "technical_skills": technical_skills,
        "soft_skills": [],
        "other_relevant_information": extract_other_information(sections, text)
    }